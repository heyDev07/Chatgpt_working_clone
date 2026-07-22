import csv
import io

import docx
import openpyxl
import pypdf

from app.core.exceptions import ValidationAppError


def _extract_pdf(data: bytes) -> str:
    reader = pypdf.PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs if p.text.strip())


def _extract_txt(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def _extract_csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [", ".join(cell.strip() for cell in row) for row in reader]
    return "\n".join(rows)


def _extract_xlsx(data: bytes) -> str:
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sheets_text = []
    for sheet in workbook.worksheets:
        rows_text = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                rows_text.append(", ".join(cells))
        if rows_text:
            sheets_text.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows_text))
    return "\n\n".join(sheets_text)


_EXTRACTORS = {
    "application/pdf": _extract_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _extract_docx,
    "text/plain": _extract_txt,
    "text/csv": _extract_csv,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": _extract_xlsx,
}


def extract_text(content_type: str, data: bytes) -> str:
    extractor = _EXTRACTORS.get(content_type)
    if not extractor:
        raise ValidationAppError(f"No text extractor available for '{content_type}'")
    try:
        text = extractor(data)
    except Exception as exc:
        raise ValidationAppError(f"Failed to extract text from file: {exc}") from exc
    if not text.strip():
        raise ValidationAppError("No extractable text found in file")
    return text
